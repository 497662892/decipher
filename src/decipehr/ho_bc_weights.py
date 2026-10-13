"""Import Ho 2022's published ASN/EUR components; never estimate SNP effects."""

from __future__ import annotations

import csv
import gzip
import math
import re
from collections import Counter
from datetime import datetime, timezone
from importlib.metadata import version
from pathlib import Path

import yaml

from decipehr.external_pgs import (
    load_weight_lock,
    read_weight_header,
    reviewed_resources,
    sha256_file,
    write_json,
)

ALLELES = set("ACGT")
COMPLEMENT = str.maketrans("ACGT", "TGCA")
COLUMNS = [
    "rsID", "chr_name", "chr_position", "effect_allele", "other_allele", "effect_weight",
    "hm_source", "hm_rsID", "hm_chr", "hm_pos",
]


def lift_snv(chrom: int, position: int, effect: str, other: str, converter) -> tuple:
    """Use pyliftover; reject ambiguous mappings and complement reverse-strand alleles."""
    hits = converter.convert_coordinate(f"chr{chrom}", position - 1)
    if not hits:
        return "", "", effect, other, "unmapped"
    if len(hits) != 1:
        return "", "", effect, other, "multiple_mappings"
    target, zero_position, strand, _ = hits[0]
    if target != f"chr{chrom}" or zero_position < 0 or strand not in {"+", "-"}:
        return "", "", effect, other, "noncanonical_mapping"
    if strand == "-":
        effect, other = effect.translate(COMPLEMENT), other.translate(COMPLEMENT)
    return str(chrom), zero_position + 1, effect, other, f"mapped_{strand}"


def convert_workbook(source: Path, destination: Path, *, component: str, score_id: str,
                     expected_count: int, converter) -> dict:
    """Preserve every published row, including unmapped rows in the coverage denominator."""
    from openpyxl import load_workbook

    expected_header = [
        "CHR", "SNV", "Position1", "Eff. Allele2", "Ref. allele3", f"PRSGW_{component}4",
    ]
    book = load_workbook(source, read_only=True, data_only=True)
    counts: Counter = Counter()
    seen = set()
    temporary = destination.with_suffix(destination.suffix + ".part")
    try:
        if {f"Chr{i}" for i in range(1, 23)} - set(book.sheetnames):
            raise ValueError("Published workbook must contain all 22 autosomes")
        description = " ".join(
            str(cell) for sheet in book if not sheet.title.startswith("Chr")
            for row in sheet.iter_rows(values_only=True) for cell in row if cell is not None
        )
        if "Build 37 position" not in description or f"PRSGW_{component}" not in description:
            raise ValueError("Workbook does not identify the expected build and component")
        with gzip.open(temporary, "wt", encoding="utf-8", newline="") as stream:
            stream.write(
                f"#score_id={score_id}\n#pgs_name=Ho2022_PRSGW_{component}\n"
                "#trait_reported=Breast cancer\n#genome_build=GRCh38\n"
                "#source_genome_build=GRCh37\n#weight_type=beta\n"
                f"#variants_number={expected_count}\n#HmPOS_build=GRCh38\n"
                f"#HmPOS_date={datetime.now(timezone.utc).date().isoformat()}\n"
                "#citation=Ho et al. Genetics in Medicine (2022); doi:10.1016/j.gim.2021.11.008\n"
            )
            writer = csv.writer(stream, delimiter="\t", lineterminator="\n")
            writer.writerow(COLUMNS)
            for chrom in range(1, 23):
                rows = book[f"Chr{chrom}"].iter_rows(values_only=True)
                if list(next(rows)) != expected_header:
                    raise ValueError(f"Unexpected published columns in Chr{chrom}")
                for row in rows:
                    if all(value is None for value in row):
                        counts["empty_formatting_rows"] += 1
                        continue
                    c, rsid, position, effect, other, beta = row
                    if (c != chrom or not isinstance(position, (int, float))
                            or not math.isfinite(position) or position < 1 or int(position) != position
                            or not isinstance(rsid, str) or not re.fullmatch(r"rs\d+", rsid)
                            or effect not in ALLELES or other not in ALLELES or effect == other
                            or not isinstance(beta, (int, float)) or not math.isfinite(beta)):
                        raise ValueError(f"Invalid published variant row in Chr{chrom}")
                    if rsid in seen:
                        raise ValueError(f"Duplicate published rsID: {rsid}")
                    seen.add(rsid)
                    new_chrom, new_pos, new_effect, new_other, status = lift_snv(
                        chrom, int(position), effect, other, converter,
                    )
                    # Strand complementation changes the allele spelling, never the beta sign.
                    writer.writerow([
                        rsid, new_chrom, new_pos, new_effect, new_other, repr(float(beta)),
                        "liftover" if new_chrom else "Unknown", rsid, new_chrom, new_pos,
                    ])
                    counts[status] += 1
                    counts["published_variants"] += 1
        if counts["published_variants"] != expected_count:
            raise ValueError("Published variant count differs from the frozen source audit")
        temporary.replace(destination)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    return dict(counts)


def prepare_ho_weights(manifest_path: str | Path, chain_path: str | Path,
                       out_dir: str | Path) -> dict:
    """Convert frozen supplementary workbooks and create a reviewed external-weight lock."""
    from pyliftover import LiftOver

    manifest_path, chain_path, out = map(Path, (manifest_path, chain_path, out_dir))
    manifest = yaml.safe_load(manifest_path.read_text())
    if manifest.get("source_build") != "GRCh37" or manifest.get("build") != "GRCh38":
        raise ValueError("Ho 2022 importer only supports GRCh37 to GRCh38")
    if sha256_file(chain_path) != manifest["chain_sha256"]:
        raise ValueError("UCSC chain differs from the frozen manifest")
    selected = reviewed_resources(manifest, {
        "target_cohort": manifest["target_cohort"], "use_context": manifest["use_context"],
        "scores": {r["id"]: r["review"] for r in manifest["resources"]},
    })
    for item in selected:
        if sha256_file(item["source_path"]) != item["source_sha256"]:
            raise ValueError(f"Published workbook has changed: {item['id']}")
    out.mkdir(parents=True, exist_ok=False)
    converter = LiftOver(str(chain_path))
    resources, qc = [], []
    for item in selected:
        path = out / f"{item['id']}_hmPOS_GRCh38.txt.gz"
        stats = convert_workbook(
            Path(item["source_path"]), path, component=item["component"], score_id=item["id"],
            expected_count=item["published_variants"], converter=converter,
        )
        resources.append({
            **item, "path": str(path.resolve()), "sha256": sha256_file(path),
            "source_path": str(Path(item["source_path"]).resolve()),
            "chain_path": str(chain_path.resolve()), "chain_sha256": sha256_file(chain_path),
            "header": read_weight_header(path, accession=item["id"], build="GRCh38"),
        })
        qc.append({"id": item["id"], **stats})
        print(f"Converted {item['id']}: {stats}", flush=True)
    lock = {
        "schema_version": 1, "build": "GRCh38", "target_cohort": manifest["target_cohort"],
        "use_context": manifest["use_context"], "resources": resources,
        "manifest_path": str(manifest_path.resolve()), "manifest_sha256": sha256_file(manifest_path),
        "participant_linkage_performed": False,
    }
    write_json(out / "weights.lock.json", lock)
    load_weight_lock(out / "weights.lock.json")
    record = {
        "status": "weights_prepared_not_scored", "snp_effects_refitted": False,
        "beta_values_changed": False, "source_build": "GRCh37", "build": "GRCh38",
        "scoring_coverage_measured": False, "patient_scoring_run": False,
        "liftover_qc": qc, "weights_lock_sha256": sha256_file(out / "weights.lock.json"),
        "versions": {name: version(name) for name in ("openpyxl", "pyliftover")},
    }
    write_json(out / "PREPARED.json", record)
    return record
